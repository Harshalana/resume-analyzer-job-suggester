"""
CareerCraft AI - Resume Parsing & Extraction Engine
Extracts structured data from PDF, DOCX, and raw text resumes.
"""

import io
import re
from typing import Dict, List, Any, Optional
from pypdf import PdfReader
import docx

from backend.taxonomy import SKILL_TAXONOMY, CANONICAL_SKILL_MAP, ALL_ACTION_VERBS


def extract_text_from_pdf(pdf_bytes: bytes) -> str:
    """Extract clean text from PDF binary content."""
    try:
        reader = PdfReader(io.BytesIO(pdf_bytes))
        extracted_text = []
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                extracted_text.append(page_text)
        return "\n".join(extracted_text)
    except Exception as e:
        raise ValueError(f"Failed to extract text from PDF: {str(e)}")


def extract_text_from_docx(docx_bytes: bytes) -> str:
    """Extract clean text from DOCX binary content."""
    try:
        doc = docx.Document(io.BytesIO(docx_bytes))
        extracted_text = [paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    extracted_text.append(row_text)
        return "\n".join(extracted_text)
    except Exception as e:
        raise ValueError(f"Failed to extract text from DOCX: {str(e)}")


def extract_contact_info(text: str) -> Dict[str, Optional[str]]:
    """Extract contact information (email, phone, urls, github, linkedin, name)."""
    # Email regex
    email_match = re.search(r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b', text)
    email = email_match.group(0) if email_match else None

    # Phone regex
    phone_match = re.search(r'(?:(?:\+?1\s*(?:[.-]\s*)?)?(?:\(\s*([2-9]1[02-9]|[2-9][02-8]1|[2-9][02-8][02-9])\s*\)|([2-9]1[02-9]|[2-9][02-8]1|[2-9][02-8][02-9]))\s*(?:[.-]\s*)?)?([2-9]1[02-9]|[2-9][02-9]1|[2-9][02-9]{2})\s*(?:[.-]\s*)?([0-9]{4})(?:\s*(?:#|x\.?|ext\.?|extension)\s*(\d+))?', text)
    phone = phone_match.group(0).strip() if phone_match and len(phone_match.group(0).strip()) >= 10 else None

    # LinkedIn
    linkedin_match = re.search(r'(?:https?://)?(?:www\.)?linkedin\.com/in/([a-zA-Z0-9_-]+)', text, re.IGNORECASE)
    linkedin = f"https://linkedin.com/in/{linkedin_match.group(1)}" if linkedin_match else None

    # GitHub
    github_match = re.search(r'(?:https?://)?(?:www\.)?github\.com/([a-zA-Z0-9_-]+)', text, re.IGNORECASE)
    github = f"https://github.com/{github_match.group(1)}" if github_match else None

    # Portfolio / Website
    portfolio_match = re.search(r'(?:https?://)?(?:www\.)?([a-zA-Z0-9-]+\.(?:dev|io|me|com|net|org)(?:/[^\s]*)?)', text, re.IGNORECASE)
    portfolio = portfolio_match.group(0) if portfolio_match and "linkedin" not in portfolio_match.group(0) and "github" not in portfolio_match.group(0) else None

    # Candidate Name heuristic (first 1-3 non-empty lines usually contains the name)
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    candidate_name = "Candidate"
    if lines:
        first_line = lines[0]
        # If first line is not too long and doesn't have an email or URL, treat it as candidate name
        if len(first_line.split()) <= 4 and "@" not in first_line and "http" not in first_line and len(first_line) < 40:
            candidate_name = first_line
        elif len(lines) > 1 and len(lines[1].split()) <= 4 and "@" not in lines[1] and len(lines[1]) < 40:
            candidate_name = lines[1]

    return {
        "name": candidate_name,
        "email": email,
        "phone": phone,
        "linkedin": linkedin,
        "github": github,
        "portfolio": portfolio
    }


def extract_skills(text: str) -> Dict[str, Any]:
    """Scan resume text against canonical skill taxonomy and categorize matches."""
    text_lower = " " + text.lower() + " "
    # Replace common separators with spaces for boundary detection
    clean_text = re.sub(r'[,/|;()\[\]•*]', ' ', text_lower)
    
    categorized_skills: Dict[str, List[str]] = {}
    all_found_skills: set = set()

    for category, skill_list in SKILL_TAXONOMY.items():
        matched_in_cat = []
        for skill in skill_list:
            # Word boundary check for accurate matching (e.g. avoid 'c' in 'cat', or 'go' in 'good')
            if skill in ["c", "r", "go", "ci cd", "ci/cd"]:
                pattern = r'(?:\b|\s)' + re.escape(skill) + r'(?:\b|\s)'
            else:
                pattern = r'\b' + re.escape(skill) + r'\b'
                
            if re.search(pattern, clean_text):
                canonical = CANONICAL_SKILL_MAP.get(skill, skill)
                formatted_name = skill.title() if len(skill) > 3 and not skill.startswith("c#") and not skill.startswith("c++") else skill.upper()
                if skill.lower() in ["react", "react.js", "next.js", "vue.js", "node.js", "fastapi", "django", "flask", "docker", "kubernetes", "terraform", "postgresql", "mysql", "mongodb", "redis", "aws", "gcp", "azure", "graphql", "tailwind css", "typescript", "javascript", "python", "pytorch", "tensorflow"]:
                    formatted_name = canonical.capitalize() if "." not in canonical else canonical
                
                if formatted_name not in matched_in_cat:
                    matched_in_cat.append(formatted_name)
                    all_found_skills.add(canonical)
                    
        if matched_in_cat:
            categorized_skills[category] = sorted(matched_in_cat)

    return {
        "categorized": categorized_skills,
        "all_skills": sorted(list(all_found_skills)),
        "total_count": len(all_found_skills)
    }


def split_sections(text: str) -> Dict[str, str]:
    """Break resume down into recognized standard sections."""
    section_headers = {
        "summary": ["summary", "professional summary", "about", "about me", "profile", "objective", "career objective"],
        "experience": ["experience", "work experience", "professional experience", "employment history", "work history", "history"],
        "education": ["education", "academic background", "academic qualifications", "academics", "degrees"],
        "skills": ["skills", "technical skills", "core competencies", "competencies", "technologies", "tech stack", "tools & technologies"],
        "projects": ["projects", "personal projects", "key projects", "notable projects", "academic projects", "open source"],
        "certifications": ["certifications", "certificates", "licenses", "courses", "professional certifications"],
        "achievements": ["achievements", "awards", "honors", "publications", "patents"]
    }

    # Find section positions
    lines = text.splitlines()
    found_sections: List[Dict[str, Any]] = []

    for i, line in enumerate(lines):
        clean_line = line.strip().lower()
        # Clean markdown headers or symbols
        clean_line = re.sub(r'^[#*\-_\s:]+|[#*\-_\s:]+$', '', clean_line).strip()
        
        for section_key, aliases in section_headers.items():
            if clean_line in aliases:
                found_sections.append({
                    "key": section_key,
                    "title": line.strip(),
                    "line_idx": i
                })
                break

    # Extract text slices between section headers
    sections_dict: Dict[str, str] = {}
    if not found_sections:
        # Fallback if no explicit headers found
        sections_dict["general"] = text
        return sections_dict

    # Sort by line index
    found_sections.sort(key=lambda x: x["line_idx"])

    # Header content before first section
    if found_sections[0]["line_idx"] > 0:
        sections_dict["header"] = "\n".join(lines[:found_sections[0]["line_idx"]]).strip()

    for idx, sec in enumerate(found_sections):
        start_line = sec["line_idx"] + 1
        end_line = found_sections[idx + 1]["line_idx"] if idx + 1 < len(found_sections) else len(lines)
        section_content = "\n".join(lines[start_line:end_line]).strip()
        sections_dict[sec["key"]] = section_content

    return sections_dict


def extract_bullet_points(text: str) -> List[Dict[str, Any]]:
    """Extract individual bullet points and analyze whether they have action verbs and metrics."""
    bullet_regex = r'(?:^|\n)\s*(?:[•\-\*–—\d+\.]\s+)(.+?)(?=(?:\n\s*[•\-\*–—\d+\.]\s+)|\Z)'
    matches = re.findall(bullet_regex, text, re.DOTALL)
    
    # If no bullet symbols found, split by lines that look like statements
    if not matches:
        matches = [line.strip() for line in text.splitlines() if len(line.strip().split()) >= 4]

    analyzed_bullets = []
    
    # Metric pattern: percentages (25%), dollar amounts ($50k), multipliers (3x), or numbers with context
    metric_pattern = re.compile(r'(\b\d+(?:\.\d+)?%|\$\s*\d+(?:,\d+)*(?:\.\d+)?[kKmMbB]?|\b\d+x\b|\b\d+(?:,\d+)*\+?\s*(?:users|clients|requests|ms|seconds|minutes|hours|days|engineers|team members|endpoints|cost|revenue|growth|reduction|transactions|queries|rps))', re.IGNORECASE)

    for bullet in matches:
        clean_b = bullet.strip().replace('\n', ' ')
        if len(clean_b) < 15 or len(clean_b.split()) < 3:
            continue
            
        first_word = clean_b.split()[0].lower().strip('.,:;()[]')
        has_action_verb = first_word in ALL_ACTION_VERBS
        
        # Check if any action verb is near the start
        has_any_action_verb = any(word in ALL_ACTION_VERBS for word in [w.lower().strip('.,:;') for w in clean_b.split()[:3]])
        
        metrics_found = metric_pattern.findall(clean_b)
        has_metrics = len(metrics_found) > 0

        analyzed_bullets.append({
            "text": clean_b,
            "has_action_verb": has_action_verb or has_any_action_verb,
            "has_metrics": has_metrics,
            "metrics": metrics_found,
            "word_count": len(clean_b.split()),
            "score": (50 if (has_action_verb or has_any_action_verb) else 0) + (50 if has_metrics else 0)
        })

    return analyzed_bullets


def parse_resume(raw_text: str) -> Dict[str, Any]:
    """Master resume parsing pipeline."""
    clean_text = raw_text.strip()
    words = clean_text.split()
    
    contacts = extract_contact_info(clean_text)
    sections = split_sections(clean_text)
    skills_data = extract_skills(clean_text)
    
    # Extract bullets from experience/projects or full text
    exp_text = sections.get("experience", "") + "\n" + sections.get("projects", "")
    bullets = extract_bullet_points(exp_text if len(exp_text.strip()) > 50 else clean_text)
    
    # Action verbs present in entire resume
    found_action_verbs = []
    text_words_lower = [re.sub(r'[^a-zA-Z]', '', w).lower() for w in words]
    for verb in ALL_ACTION_VERBS:
        if verb in text_words_lower:
            found_action_verbs.append(verb)

    return {
        "raw_text": clean_text,
        "word_count": len(words),
        "char_count": len(clean_text),
        "contact": contacts,
        "sections": sections,
        "skills": skills_data,
        "bullets": bullets,
        "action_verbs": sorted(list(set(found_action_verbs)))
    }

